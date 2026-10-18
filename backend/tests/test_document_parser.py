"""文档解析 / 清洗 / 切片 / ingest_file 单元测试。

所有 PDF / DOCX / XLSX 测试文件均在运行时程序化生成到临时目录，
不依赖任何外部真实文档或大文件。

在 backend/ 目录下运行：
    python -m pytest tests/test_document_parser.py -v
"""

from __future__ import annotations

import pytest
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import settings
from services.document_parser import (
    CorruptedDocumentError,
    DocumentFragment,
    EmptyDocumentError,
    EncryptedDocumentError,
    UnsupportedFileTypeError,
    clean_text,
    parse_document,
    remove_repeated_header_footer,
)
from services.ingest import FileTooLargeError, ingest_file, validate_upload_file
from services.text_splitter import build_splitter, split_fragments


# ---------------------------------------------------------------------------
# 测试文件构造工具
# ---------------------------------------------------------------------------

def _make_pdf(path, pages: int = 3, encrypted: bool = False, empty: bool = False) -> None:
    """生成含中文的多页 PDF：重复页眉 + 页脚页码 + 每页不同正文。"""
    import fitz

    doc = fitz.open()
    if not empty:
        for i in range(1, pages + 1):
            page = doc.new_page()
            # china-s 为 MuPDF 内置简体中文字体，写入后可被 get_text 提取
            page.insert_text((72, 40), "公司内部手册", fontname="china-s", fontsize=10)
            page.insert_text(
                (72, 110), f"这是第{i}页的正文内容，包含制度条款。",
                fontname="china-s", fontsize=12,
            )
            page.insert_text(
                (72, 140), f"第{i}页需要能被完整提取出来用于检索。",
                fontname="china-s", fontsize=12,
            )
            page.insert_text((280, 760), f"第 {i} 页", fontname="china-s", fontsize=10)
    else:
        doc.new_page()

    if encrypted:
        doc.save(
            str(path),
            encryption=fitz.PDF_ENCRYPT_AES_256,
            owner_pw="owner-secret",
            user_pw="user-secret",
        )
    else:
        doc.save(str(path))
    doc.close()


def _make_docx(path) -> None:
    """生成含两个标题章节、段落、表格的 DOCX。"""
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 总则", level=1)
    doc.add_paragraph("第一条 为规范内部管理，特制定本制度。")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "姓名"
    table.rows[0].cells[1].text = "部门"
    table.rows[1].cells[0].text = "张三"
    table.rows[1].cells[1].text = "研发部"
    doc.add_heading("第二章 附则", level=1)
    doc.add_paragraph("本制度自发布之日起施行。")
    doc.save(str(path))


def _make_xlsx(path, with_formula: bool = False) -> None:
    """生成含两个工作表的 XLSX。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "人员表"
    ws1.append(["姓名", "部门"])
    ws1.append(["张三", "研发部"])
    if with_formula:
        ws1.append(["合计", "=1+2"])  # data_only 读取时该格为 None（从未被 Excel 计算）

    ws2 = wb.create_sheet("资产表")
    ws2.append(["资产", "数量"])
    ws2.append(["笔记本电脑", 10])
    wb.save(str(path))


def _write_garbage(path, content: bytes = b"this is not a real office document") -> None:
    path.write_bytes(content)


# ---------------------------------------------------------------------------
# 文本清洗
# ---------------------------------------------------------------------------

def test_clean_text_removes_control_chars_and_noise() -> None:
    raw = "正文第一行\x00\x07　　多余空格\n第 3 页\n\n\n正文第二行"
    result = clean_text(raw)
    assert "\x00" not in result and "\x07" not in result
    assert "第 3 页" not in result          # 页码噪声行被过滤
    assert "正文第一行" in result
    assert "正文第二行" in result
    assert "　　" not in result              # 全角空白被压缩
    assert "\n\n\n" not in result            # 连续空行被压缩


def test_clean_text_handles_empty_and_deduplicates() -> None:
    assert clean_text("") == ""
    assert clean_text("相同行\n相同行\n不同行") == "相同行\n不同行"
    # 英文页码噪声
    assert "Page 5 of 10" not in clean_text("标题\nPage 5 of 10\n正文")


def test_remove_repeated_header_footer_across_pages() -> None:
    pages = [
        "公司机密\n第一页独有内容A\n1",
        "公司机密\n第二页独有内容B\n2",
        "公司机密\n第三页独有内容C\n3",
    ]
    result = remove_repeated_header_footer(pages)
    assert all("公司机密" not in text for text in result)
    assert "第一页独有内容A" in result[0]
    # 不足 3 页时不做跨页过滤
    short = remove_repeated_header_footer(["公司机密\n内容一", "公司机密\n内容二"])
    assert "公司机密" in short[0]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_parse_pdf_preserves_real_pages_and_filters_noise(tmp_path) -> None:
    pdf_path = tmp_path / "制度手册.pdf"
    _make_pdf(pdf_path)

    fragments = parse_document(pdf_path)

    assert len(fragments) == 3
    assert [f.page for f in fragments] == [1, 2, 3]  # 真实页码，从 1 开始
    for fragment in fragments:
        assert fragment.file_type == "pdf"
        assert fragment.doc_name == "制度手册.pdf"
        assert fragment.section is None              # PDF 不使用 section
        assert fragment.source_path.endswith("制度手册.pdf")
        assert "公司内部手册" not in fragment.text    # 跨页重复页眉已去除
        assert "第 " not in fragment.text            # 页脚页码已去除
        assert "正文" in fragment.text


def test_parse_pdf_encrypted_raises_readable_error(tmp_path) -> None:
    pdf_path = tmp_path / "加密文件.pdf"
    _make_pdf(pdf_path, encrypted=True)

    with pytest.raises(EncryptedDocumentError) as exc_info:
        parse_document(pdf_path)
    assert "加密" in str(exc_info.value)


def test_parse_pdf_corrupt_raises(tmp_path) -> None:
    pdf_path = tmp_path / "损坏.pdf"
    _write_garbage(pdf_path, b"%PDF-broken garbage data")

    with pytest.raises(CorruptedDocumentError):
        parse_document(pdf_path)


def test_parse_pdf_empty_raises(tmp_path) -> None:
    pdf_path = tmp_path / "空白.pdf"
    _make_pdf(pdf_path, empty=True)

    with pytest.raises(EmptyDocumentError) as exc_info:
        parse_document(pdf_path)
    assert "扫描件" in str(exc_info.value)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_parse_docx_tracks_sections_and_table(tmp_path) -> None:
    docx_path = tmp_path / "管理制度.docx"
    _make_docx(docx_path)

    fragments = parse_document(docx_path)

    assert len(fragments) == 2
    assert [f.section for f in fragments] == ["第一章 总则", "第二章 附则"]
    for fragment in fragments:
        assert fragment.page is None               # 禁止伪造页码
        assert fragment.file_type == "docx"
    # 第一章片段包含段落正文与表格内容
    chapter1 = fragments[0].text
    assert "第一条" in chapter1
    assert "张三" in chapter1 and "研发部" in chapter1
    assert " | " in chapter1                        # 表格行被序列化
    assert "第二章 附则" in fragments[1].text


def test_parse_docx_corrupt_raises(tmp_path) -> None:
    docx_path = tmp_path / "损坏.docx"
    _write_garbage(docx_path)

    with pytest.raises(CorruptedDocumentError) as exc_info:
        parse_document(docx_path)
    message = str(exc_info.value)
    assert ".docx" in message or "加密" in message


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def test_parse_xlsx_uses_sheet_name_as_section(tmp_path) -> None:
    xlsx_path = tmp_path / "台账.xlsx"
    _make_xlsx(xlsx_path)

    fragments = parse_document(xlsx_path)

    assert len(fragments) == 2
    assert [f.section for f in fragments] == ["人员表", "资产表"]
    assert fragments[0].page is None                # Excel 无页码
    assert fragments[0].file_type == "xlsx"
    assert "张三" in fragments[0].text and "研发部" in fragments[0].text
    assert "笔记本电脑" in fragments[1].text and "10" in fragments[1].text


def test_parse_xlsx_data_only_skips_uncalculated_formula(tmp_path) -> None:
    xlsx_path = tmp_path / "含公式.xlsx"
    _make_xlsx(xlsx_path, with_formula=True)

    fragments = parse_document(xlsx_path)
    # data_only=True 取缓存计算结果；openpyxl 写入的公式从未被 Excel 计算 → None
    assert "=1+2" not in fragments[0].text
    assert "张三" in fragments[0].text


def test_parse_xlsx_corrupt_raises(tmp_path) -> None:
    xlsx_path = tmp_path / "损坏.xlsx"
    _write_garbage(xlsx_path)

    with pytest.raises(CorruptedDocumentError):
        parse_document(xlsx_path)


# ---------------------------------------------------------------------------
# 统一入口 / 元数据 / 切片
# ---------------------------------------------------------------------------

def test_parse_document_unsupported_type(tmp_path) -> None:
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("hello", encoding="utf-8")

    with pytest.raises(UnsupportedFileTypeError) as exc_info:
        parse_document(txt_path)
    assert "pdf" in str(exc_info.value)  # 错误信息提示支持的类型


def test_fragment_metadata_aligns_with_schema(tmp_path) -> None:
    pdf_path = tmp_path / "a.pdf"
    _make_pdf(pdf_path, pages=1)
    fragment = parse_document(pdf_path)[0]

    meta = fragment.metadata()
    assert set(meta.keys()) == {
        "doc_name", "source_path", "file_type", "page", "section", "chunk_index"
    }
    assert meta["page"] == 1 and meta["chunk_index"] is None


def test_split_fragments_preserves_metadata_and_indexes(tmp_path) -> None:
    docx_path = tmp_path / "长文档.docx"
    from docx import Document

    doc = Document()
    doc.add_heading("章节甲", level=1)
    # 足够长的正文以产生多个切片
    doc.add_paragraph("这是一段用于测试切片的中文正文。" * 60)
    doc.save(str(docx_path))
    fragments = parse_document(docx_path)

    tiny_splitter = RecursiveCharacterTextSplitter(
        chunk_size=80, chunk_overlap=10,
        separators=["\n", "。", "，", " ", ""],
    )
    chunks = split_fragments(fragments, splitter=tiny_splitter)

    assert len(chunks) > 1
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    for chunk in chunks:
        assert len(chunk.text) <= 80
        assert chunk.section == "章节甲"
        assert chunk.page is None
        assert chunk.doc_name == "长文档.docx"


def test_build_splitter_reads_settings_and_validates(monkeypatch) -> None:
    splitter = build_splitter()
    assert splitter._chunk_size == settings.chunk_size
    assert splitter._chunk_overlap == settings.chunk_overlap

    monkeypatch.setattr(settings, "chunk_overlap", settings.chunk_size + 1)
    with pytest.raises(ValueError):
        build_splitter()


# ---------------------------------------------------------------------------
# ingest_file 上层入口
# ---------------------------------------------------------------------------

def test_ingest_file_happy_path(tmp_path) -> None:
    docx_path = tmp_path / "入库测试.docx"
    _make_docx(docx_path)

    chunks = ingest_file(docx_path)
    assert len(chunks) >= 2
    assert all(c.chunk_index is not None for c in chunks)
    assert all(isinstance(c, DocumentFragment) for c in chunks)


def test_ingest_file_rejects_unsupported_extension(tmp_path) -> None:
    txt_path = tmp_path / "x.txt"
    txt_path.write_text("abc", encoding="utf-8")
    with pytest.raises(UnsupportedFileTypeError):
        ingest_file(txt_path)


def test_ingest_file_rejects_oversize(tmp_path, monkeypatch) -> None:
    xlsx_path = tmp_path / "big.xlsx"
    _make_xlsx(xlsx_path)

    # 把大小上限压到 10 字节（替换 property，避免 pydantic 字段描述符拦截）
    monkeypatch.setattr(type(settings), "max_upload_bytes", property(lambda self: 10))
    with pytest.raises(FileTooLargeError) as exc_info:
        validate_upload_file(xlsx_path)
    assert "MB" in str(exc_info.value)


def test_ingest_file_missing_file(tmp_path) -> None:
    with pytest.raises(Exception):
        ingest_file(tmp_path / "not-exist.pdf")


def test_ingest_file_propagates_encrypted_error(tmp_path) -> None:
    pdf_path = tmp_path / "secret.pdf"
    _make_pdf(pdf_path, encrypted=True)
    with pytest.raises(EncryptedDocumentError):
        ingest_file(pdf_path)
